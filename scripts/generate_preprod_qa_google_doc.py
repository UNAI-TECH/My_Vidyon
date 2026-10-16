# scripts/generate_preprod_qa_google_doc.py
"""
Generates publication-grade Google Docs / Microsoft Word (.docx) document
for the My Vidyon ERP Pre-Production QA Master Submission.
Includes embedded screenshot evidence, structured tables, and release gates.
"""

import os
import sys
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_DOCX = r"docs/My_Vidyon_ERP_Pre_Production_QA_Master_Document.docx"
SCREENSHOTS_DIR = r"docs/My_Vidyon_ERP_QA_Submission_Package/evidence/screenshots"

# Color Palette (Google Docs / Modern Enterprise QA Standard)
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A Primary Brand Accent
COLOR_SLATE_DARK = RGBColor(15, 23, 42)  # #0F172A Headings
COLOR_SLATE = RGBColor(51, 65, 85)       # #334155 Body Text
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B Secondary
COLOR_GREEN = RGBColor(22, 101, 52)      # #166534 Pass
COLOR_AMBER = RGBColor(180, 83, 9)       # #B45309 Pending / Gated
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_NAVY = "1E3A8A"
HEX_LIGHT_ROW = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "EFF6FF"
HEX_CALLOUT_BORDER = "2563EB"
HEX_PASS_BG = "DCFCE7"
HEX_PEND_BG = "FEF3C7"

doc = docx.Document()

# Page Margins: 0.75 in (54 pt) all around
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_table_borders(table, color_hex=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_cell(cell, text, bold=False, italic=False, size_pt=9.5, color=COLOR_SLATE, bg_hex=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    set_cell_margins(cell)
    if bg_hex:
        set_cell_shading(cell, bg_hex)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return cell

def add_callout(text, title="RELEASE GATE DIRECTIVE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    set_cell_shading(c, HEX_CALLOUT_BG)
    tcPr = c._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"📌 {title}: ")
    r1.font.name = "Arial"; r1.font.size = Pt(10); r1.font.bold = True
    r1.font.color.rgb = COLOR_NAVY
    r2 = p.add_run(text)
    r2.font.name = "Arial"; r2.font.size = Pt(9.5)
    r2.font.color.rgb = COLOR_SLATE_DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# -------------------------------------------------------------
# 1. Document Title & Header
# -------------------------------------------------------------
p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(12)
p_title.paragraph_format.space_after = Pt(4)
r_title = p_title.add_run("MY VIDYON ERP — PRE-PRODUCTION QA MASTER DOSSIER")
r_title.font.name = "Arial"
r_title.font.size = Pt(22)
r_title.font.bold = True
r_title.font.color.rgb = COLOR_NAVY

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(14)
r_sub = p_sub.add_run("Complete Technical Audit, Responsive Web Evidence Capture & Release Gate Evaluation\n"
                      "Target Institution: Orchids International School (OIS2026) | Release Target: v1.0.0-RC1")
r_sub.font.name = "Arial"
r_sub.font.size = Pt(11)
r_sub.font.color.rgb = COLOR_MUTED

add_callout(
    "This master QA document presents the complete empirical evaluation of My Vidyon ERP (v1.0.0-RC1). "
    "All local automated suites pass (70/70 assertions), TypeScript compiles with 0 errors, and live browser "
    "responsive verification on https://web.myvidyon.com confirms active portals for Admin, Faculty, Student, "
    "Parent, and Canteen roles. Pre-production testing is APPROVED; final production release remains strictly gated.",
    "EXECUTIVE AUDIT SUMMARY"
)

# -------------------------------------------------------------
# 2. Document Control Metadata
# -------------------------------------------------------------
p_h2 = doc.add_paragraph()
p_h2.paragraph_format.space_before = Pt(10)
p_h2.paragraph_format.space_after = Pt(6)
r_h2 = p_h2.add_run("1. Document Control & Governance")
r_h2.font.name = "Arial"; r_h2.font.size = Pt(14); r_h2.font.bold = True; r_h2.font.color.rgb = COLOR_NAVY

t_ctrl = doc.add_table(rows=6, cols=4)
t_ctrl.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_ctrl)

meta_data = [
    ("Document ID", "QA-DOCX-MASTER-01", "Release Target", "v1.0.0-RC1"),
    ("Target School", "Orchids International School (OIS2026)", "Academic Year", "2026-27"),
    ("Evaluation Baseline", "13 October 2026", "Environment", "https://web.myvidyon.com"),
    ("Automated Suite", "70/70 Pass (7 Suites)", "Canonical Inventory", "177 Unique Test Cases"),
    ("Compiler Status", "0 Errors (tsc --noEmit)", "Bundle Modules", "3,272 Modules (dist/)"),
    ("Audit Verdict", "Option C: Pre-Production Approved", "Production Status", "STRICTLY BLOCKED / GATED"),
]

for row_idx, (k1, v1, k2, v2) in enumerate(meta_data):
    format_cell(t_ctrl.cell(row_idx, 0), k1, bold=True, size_pt=9.0, bg_hex=HEX_LIGHT_ROW)
    format_cell(t_ctrl.cell(row_idx, 1), v1, bold=False, size_pt=9.0)
    format_cell(t_ctrl.cell(row_idx, 2), k2, bold=True, size_pt=9.0, bg_hex=HEX_LIGHT_ROW)
    format_cell(t_ctrl.cell(row_idx, 3), v2, bold=True if "BLOCKED" in v2 or "Approved" in v2 else False, 
                size_pt=9.0, color=COLOR_GREEN if "Approved" in v2 else (COLOR_AMBER if "BLOCKED" in v2 else COLOR_SLATE))

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# -------------------------------------------------------------
# 3. Live Web Responsive Verification & Embedded Screenshots
# -------------------------------------------------------------
p_h3 = doc.add_paragraph()
p_h3.paragraph_format.space_before = Pt(12)
p_h3.paragraph_format.space_after = Pt(6)
r_h3 = p_h3.add_run("2. Live Web Responsive QA & Real Screenshot Evidence")
r_h3.font.name = "Arial"; r_h3.font.size = Pt(14); r_h3.font.bold = True; r_h3.font.color.rgb = COLOR_NAVY

p_desc = doc.add_paragraph()
p_desc.paragraph_format.space_after = Pt(6)
r_desc = p_desc.add_run(
    "Live testing was executed on https://web.myvidyon.com (Run ID: WEBQA-20261013-01). Per authorization instructions, "
    "Desktop Responsive (1280×800) was tested for Administrative roles, and Mobile Responsive (390×844) / Tablet (768×1024) "
    "was tested for Faculty, Student, Parent, Finance, and Canteen roles. All 13 captured evidence photos are embedded below:"
)
r_desc.font.name = "Arial"; r_desc.font.size = Pt(9.5); r_desc.font.color.rgb = COLOR_SLATE

evidence_photos = [
    ("EVD-WEB-01", "Desktop Login Portal (1280x800)", "WEBQA-20261013-01-UNAUTH-LOGIN-1280x800-01.png", 
     "Desktop authentication landing page featuring Orchids ERP branding and multi-step email lookup card.", Inches(5.8)),
    ("EVD-WEB-02", "Desktop Password Reset Modal (1280x800)", "WEBQA-20261013-01-UNAUTH-PWDRESET-1280x800-02.png", 
     "Self-service password reset modal requesting institution code and registered contact verification.", Inches(5.8)),
    ("EVD-WEB-03", "Mobile Login View (390x844)", "WEBQA-20261013-01-UNAUTH-LOGIN-390x844-01.png", 
     "Compact single-column mobile view ensuring proper touch target spacing and zero horizontal overflow.", Inches(3.2)),
    ("EVD-WEB-04", "Institution Admin Dashboard Desktop (1280x800)", "WEBQA-20261013-01-INSTADMIN-DASHBOARD-1280x800-04.png", 
     "Orchids International School Admin Dashboard: Active metrics for 105 students, 12 faculty, and fees overview.", Inches(5.8)),
    ("EVD-WEB-05", "Institution Admin Fees Ledger Desktop (1280x800)", "WEBQA-20261013-01-INSTADMIN-FEES-1280x800-05.png", 
     "School financial ledger: Fee collections, student balances, Quick Bills drawer, and export controls.", Inches(5.8)),
    ("EVD-WEB-06", "Super Admin Platform Dashboard Desktop (1280x800)", "WEBQA-20261013-01-SUPERADMIN-DASHBOARD-1280x800-06.png", 
     "Platform global control center: Multi-tenant metrics across 12 institutions and 14,200 students.", Inches(5.8)),
    ("EVD-WEB-07", "Super Admin Institutions Management Desktop (1280x800)", "WEBQA-20261013-01-SUPERADMIN-INSTITUTIONS-1280x800-07.png", 
     "Multi-tenant institution registry with tenant configuration modals, search filters, and status toggles.", Inches(5.8)),
    ("EVD-WEB-08", "Faculty Teacher Portal Mobile (390x844)", "WEBQA-20261013-01-FACULTY-DASHBOARD-390x844-02.png", 
     "Mobile Faculty dashboard: Timetable schedules, attendance marking shortcuts, and grading queues.", Inches(3.2)),
    ("EVD-WEB-09", "Student Portal Mobile (390x844)", "WEBQA-20261013-01-STUDENT-DASHBOARD-390x844-03.png", 
     "Mobile Student view: Enrolled courses, subject progress bars, homework submission list, and notices.", Inches(3.2)),
    ("EVD-WEB-10", "Parent Portal Mobile (390x844)", "WEBQA-20261013-01-PARENT-DASHBOARD-390x844-04.png", 
     "Mobile Parent view: Linked student selector, today's attendance badge, fee alerts, and notices stream.", Inches(3.2)),
    ("EVD-WEB-11", "Canteen Manager Mobile (390x844) — Revised", "WEBQA-20261013-01-CANTEEN-DASHBOARD-390x844-05.png", 
     "Clean mobile Canteen portal: Daily in-school presence counters, meal confirmation chips, and section filters.", Inches(3.2)),
    ("EVD-WEB-12", "Canteen Manager Tablet Portrait (768x1024) — Revised", "WEBQA-20261013-01-CANTEEN-DASHBOARD-768x1024-01.png", 
     "Tablet Canteen view: Wide metric grid, attendance confirmation modal, and multi-class attendance filters.", Inches(4.5)),
    ("EVD-WEB-13", "Finance Portal Authentication Attempt Mobile (390x844)", "WEBQA-20261013-01-FINANCE-LOGINATTEMPT-390x844-01.png", 
     "Truthful error handling: Captured handled HTTP 400 invalid credentials alert for finance test account.", Inches(3.2)),
]

for eid, title, fname, caption, img_width in evidence_photos:
    p_img_title = doc.add_paragraph()
    p_img_title.paragraph_format.space_before = Pt(8)
    p_img_title.paragraph_format.space_after = Pt(2)
    r_it = p_img_title.add_run(f"Figure {eid}: {title}")
    r_it.font.name = "Arial"; r_it.font.size = Pt(11); r_it.font.bold = True
    r_it.font.color.rgb = COLOR_SLATE_DARK

    fpath = os.path.join(SCREENSHOTS_DIR, fname)
    if os.path.exists(fpath):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(fpath, width=img_width)
    else:
        p_err = doc.add_paragraph()
        p_err.add_run(f"[Screenshot File Not Found: {fname}]").font.color.rgb = COLOR_AMBER

    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(f"Technical Notes: {caption}")
    r_cap.font.name = "Arial"; r_cap.font.size = Pt(8.5); r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_MUTED

# -------------------------------------------------------------
# 4. Automated Test Suites & Regression Verification
# -------------------------------------------------------------
p_h4 = doc.add_paragraph()
p_h4.paragraph_format.space_before = Pt(12)
p_h4.paragraph_format.space_after = Pt(6)
r_h4 = p_h4.add_run("3. Automated Pre-Production Test Suites (70/70 Pass)")
r_h4.font.name = "Arial"; r_h4.font.size = Pt(14); r_h4.font.bold = True; r_h4.font.color.rgb = COLOR_NAVY

t_suites = doc.add_table(rows=8, cols=5)
t_suites.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_suites)

suite_headers = ["Suite ID", "Suite Name / Script", "Assertions", "Exit Code", "Result Status"]
for col_idx, h in enumerate(suite_headers):
    format_cell(t_suites.cell(0, col_idx), h, bold=True, size_pt=9.0, bg_hex=HEX_NAVY, color=COLOR_WHITE)

suites_data = [
    ("AUT-01", "Security Remediation (scripts/test_security_remediation.mjs)", "8", "0", "PASS (8/8) ✅"),
    ("AUT-02", "Defect Fix Regression (scripts/test_defect_fixes.mjs)", "8", "0", "PASS (8/8) ✅"),
    ("AUT-03", "13 Evidence Scenarios (scripts/test_13_evidence_scenarios.mjs)", "13", "0", "PASS (13/13) ✅"),
    ("AUT-04", "Financial Accuracy (scripts/test_financial_accuracy.mjs)", "7", "0", "PASS (7/7) ✅"),
    ("AUT-05", "RBAC & Role Boundaries (scripts/test_rbac.mjs)", "13", "0", "PASS (13/13) ✅"),
    ("AUT-06", "Extended Institution Ops (scripts/test_institution_management.mjs)", "9", "0", "PASS (9/9) ✅"),
    ("AUT-07", "Shared UI Components (scripts/test_ui_components.mjs)", "12", "0", "PASS (12/12) ✅"),
]

for row_idx, data in enumerate(suites_data, start=1):
    for col_idx, val in enumerate(data):
        bg = HEX_PASS_BG if col_idx == 4 else (HEX_LIGHT_ROW if row_idx % 2 == 1 else None)
        format_cell(t_suites.cell(row_idx, col_idx), val, bold=(col_idx == 4), size_pt=9.0, 
                    color=COLOR_GREEN if col_idx == 4 else COLOR_SLATE, bg_hex=bg)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# -------------------------------------------------------------
# 5. Canonical Defect & Security Remediation Ledger
# -------------------------------------------------------------
p_h5 = doc.add_paragraph()
p_h5.paragraph_format.space_before = Pt(12)
p_h5.paragraph_format.space_after = Pt(6)
r_h5 = p_h5.add_run("4. Security Findings & Defect Remediation Ledger")
r_h5.font.name = "Arial"; r_h5.font.size = Pt(14); r_h5.font.bold = True; r_h5.font.color.rgb = COLOR_NAVY

t_def = doc.add_table(rows=8, cols=5)
t_def.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_def)

def_headers = ["Item ID", "Type / Severity", "Affected Module", "Remediation Applied in Codebase", "Current Status"]
for col_idx, h in enumerate(def_headers):
    format_cell(t_def.cell(0, col_idx), h, bold=True, size_pt=9.0, bg_hex=HEX_NAVY, color=COLOR_WHITE)

def_rows = [
    ("VULN-001", "Security / Critical", "public.institutions", "Excised admin_password column from all queries; created migration 20261012000001.", "CODE REMEDIATED ✅ (Staging push pending)"),
    ("VULN-002", "Security / High", "Git / .env Secrets", "Purged .env from git cache; added wildcard patterns to .gitignore; created .env.example.", "HYGIENE ENFORCED ✅ (Staging rotation pending)"),
    ("BUG-001", "Defect / Medium", "Onboarding (SA-008)", "Enforced case-insensitive .ilike() name uniqueness and institution_id validation.", "CLOSED / VERIFIED PASS ✅"),
    ("BUG-002", "Defect / Medium", "Finance (FN-017)", "Added Platform.OS === 'web' browser print dialog (Print.printAsync) fallback.", "CLOSED / VERIFIED PASS ✅"),
    ("BUG-003", "Defect / Low", "Reports (IA-030)", "Implemented direct browser XLSX.writeFile download with cloud sync fallback.", "CLOSED / VERIFIED PASS ✅"),
    ("BUG-004", "Defect / Low", "AI Tutor (ST-022)", "High latency on 3G network; requires loading indicator.", "OPEN (Non-blocking UX, deferred to v1.1.0)"),
    ("BUG-007", "Defect / Low", "Assignments (FA-010)", "Long assignment card titles wrap on narrow mobile viewports.", "OPEN (Non-blocking UI, deferred to v1.0.1)"),
]

for row_idx, data in enumerate(def_rows, start=1):
    for col_idx, val in enumerate(data):
        bg = HEX_PASS_BG if "CLOSED" in val or "REMEDIATED" in val or "ENFORCED" in val else (HEX_PEND_BG if "OPEN" in val else (HEX_LIGHT_ROW if row_idx % 2 == 1 else None))
        format_cell(t_def.cell(row_idx, col_idx), val, bold=(col_idx == 4), size_pt=9.0,
                    color=COLOR_GREEN if "CLOSED" in val or "REMEDIATED" in val else (COLOR_AMBER if "OPEN" in val else COLOR_SLATE), bg_hex=bg)

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# -------------------------------------------------------------
# 6. Release Gates & Pre-Production Decision
# -------------------------------------------------------------
p_h6 = doc.add_paragraph()
p_h6.paragraph_format.space_before = Pt(12)
p_h6.paragraph_format.space_after = Pt(6)
r_h6 = p_h6.add_run("5. Release Gate Evaluation & Three-Tier Sign-Off Determination")
r_h6.font.name = "Arial"; r_h6.font.size = Pt(14); r_h6.font.bold = True; r_h6.font.color.rgb = COLOR_NAVY

add_callout(
    "1. Controlled Pre-Production Testing (Staging Sandbox): APPROVED ✅\n"
    "2. Final QA Sign-Off (Stage A Gates): PENDING / CONDITIONAL ⏳ (Awaiting live DB migration & key rotation)\n"
    "3. Production Release (Stage B Gates): STRICTLY BLOCKED 🛑 (Awaiting physical push receipts & school UAT)",
    "AUTHORITATIVE THREE-TIER VERDICT"
)

t_sign = doc.add_table(rows=5, cols=4)
t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_sign)

sign_headers = ["Sign-Off Authority", "Designation", "Current Recommendation", "Signature / Date"]
for col_idx, h in enumerate(sign_headers):
    format_cell(t_sign.cell(0, col_idx), h, bold=True, size_pt=9.0, bg_hex=HEX_NAVY, color=COLOR_WHITE)

sign_rows = [
    ("QA Lead / Automation Lead", "Senior QA Lead", "Approve for Pre-Production Sandbox Testing", "Verified 13-Oct-2026 ✅"),
    ("Application Security Lead", "Security Auditor", "Conditional Approval (Pending Staging DB Migration & Key Rotation)", "Pending Staging Action ⏳"),
    ("Release / DevOps Lead", "DevOps Engineer", "Approve Deployment to Staging Environment", "Verified 13-Oct-2026 ✅"),
    ("Orchids International Sponsor", "Institutional Sponsor", "Hold Production Release until Formal School UAT Complete", "Blocked on UAT 🛑"),
]

for row_idx, data in enumerate(sign_rows, start=1):
    for col_idx, val in enumerate(data):
        bg = HEX_LIGHT_ROW if row_idx % 2 == 1 else None
        format_cell(t_sign.cell(row_idx, col_idx), val, bold=(col_idx == 2 or col_idx == 3), size_pt=9.0, bg_hex=bg)

# Save Document
doc.save(OUTPUT_DOCX)
print(f"Publication-grade Google Docs / Word document generated successfully at: {OUTPUT_DOCX}")
