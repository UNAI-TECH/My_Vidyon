import docx
import os

docs_to_inspect = [
    'Copy of 08_UI_UX_Design_Document.docx',
    'Copy of 11_Test_Cases.docx',
    'Copy of 12_UAT_Sign_off.docx'
]

# Check docs/ai/documentations/ and any other folders
for doc_name in docs_to_inspect:
    path = os.path.join('docs', 'ai', 'documentations', doc_name)
    if not os.path.exists(path):
        # search workspace
        found = []
        for root, dirs, files in os.walk('.'):
            if doc_name in files:
                found.append(os.path.join(root, doc_name))
        print(f"File {doc_name} search results: {found}")
        if found:
            path = found[0]
        else:
            print(f"FILE NOT FOUND ANYWHERE: {doc_name}")
            continue

    doc = docx.Document(path)
    print(f'==============================')
    print(f'FILE: {doc_name} at {path}')
    print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f'  P{i}: "{p.text}"')
    for t_idx, table in enumerate(doc.tables):
        print(f'  Table {t_idx} (rows: {len(table.rows)}, cols: {len(table.columns)}):')
        for r_idx, row in enumerate(table.rows[:6]):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            print(f'    R{r_idx}: {cells}')
