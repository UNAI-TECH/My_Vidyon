import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_PATH = r"c:\Users\kamal\OneDrive\Desktop\vidyon\My_Vidyon\docs\My_Vidyon_Testing_Report_420_Executed.xlsx"
OUTPUT_COPY = r"c:\Users\kamal\OneDrive\Desktop\vidyon\My_Vidyon\docs\My_Vidyon_QA_Test_Execution_Logs_420.xlsx"

# Theme Colors
HEX_NAVY = "1E3A8A"        # Deep Navy Header
HEX_BLUE_ACCENT = "2563EB" # Royal Blue Accent
HEX_DARK_SLATE = "0F172A"  # Dark Slate
HEX_WHITE = "FFFFFF"
HEX_ZEBRA = "F8FAFC"
HEX_BORDER = "CBD5E1"

# Fonts
FONT_TITLE = Font(name="Segoe UI", size=16, bold=True, color=HEX_NAVY)
FONT_SUBTITLE = Font(name="Segoe UI", size=10, italic=True, color="64748B")
FONT_SECTION = Font(name="Segoe UI", size=12, bold=True, color=HEX_NAVY)
FONT_HEADER = Font(name="Segoe UI", size=9.5, bold=True, color=HEX_WHITE)
FONT_DATA = Font(name="Segoe UI", size=9, color="1E293B")
FONT_DATA_BOLD = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
FONT_KPI_VAL = Font(name="Segoe UI", size=18, bold=True, color=HEX_NAVY)
FONT_KPI_LBL = Font(name="Segoe UI", size=8.5, bold=True, color="64748B")

# Fills
FILL_HEADER = PatternFill(start_color=HEX_NAVY, end_color=HEX_NAVY, fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color=HEX_ZEBRA, end_color=HEX_ZEBRA, fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_KPI_CARD = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

# Status Fills & Fonts
FILL_PASSED = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
FONT_PASSED = Font(name="Segoe UI", size=9, bold=True, color="166534")

FILL_OBS = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
FONT_OBS = Font(name="Segoe UI", size=9, bold=True, color="92400E")

FILL_DEFERRED = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
FONT_DEFERRED = Font(name="Segoe UI", size=9, bold=True, color="475569")

# Priority Fonts
FONT_P0 = Font(name="Segoe UI", size=9, bold=True, color="B91C1C")
FONT_P1 = Font(name="Segoe UI", size=9, bold=True, color="C2410C")
FONT_P2 = Font(name="Segoe UI", size=9, color="B45309")
FONT_P3 = Font(name="Segoe UI", size=9, color="475569")

# Borders
THIN_SIDE = Side(border_style="thin", color=HEX_BORDER)
MEDIUM_SIDE = Side(border_style="medium", color=HEX_NAVY)
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
HEADER_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=MEDIUM_SIDE, bottom=MEDIUM_SIDE)
KPI_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

COLUMNS = [
    ("TC ID", 13, ALIGN_CENTER),
    ("Category / Module Group", 24, ALIGN_LEFT),
    ("Sub-Module / Feature", 20, ALIGN_LEFT),
    ("Test Scenario / Description", 35, ALIGN_LEFT),
    ("Pre-Conditions & Test Data", 25, ALIGN_LEFT),
    ("Execution Steps", 40, ALIGN_LEFT),
    ("Expected Result", 32, ALIGN_LEFT),
    ("Actual Result", 32, ALIGN_LEFT),
    ("Execution Status", 16, ALIGN_CENTER),
    ("Severity / Priority", 14, ALIGN_CENTER),
    ("Test Type", 14, ALIGN_CENTER),
    ("Target Platform", 16, ALIGN_CENTER),
    ("Defect Ref", 12, ALIGN_CENTER),
    ("Execution Date", 13, ALIGN_CENTER),
    ("Tester Name", 18, ALIGN_CENTER),
    ("Remarks & Technical Evidence", 32, ALIGN_LEFT)
]

def style_header_row(ws, row_idx):
    ws.row_dimensions[row_idx].height = 28
    for c_idx, (col_name, width, default_align) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row_idx, column=c_idx)
        cell.value = col_name
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = HEADER_BORDER
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = width

def apply_test_row(ws, row_idx, tc_data, is_zebra=False):
    ws.row_dimensions[row_idx].height = 24
    base_fill = FILL_ZEBRA if is_zebra else FILL_WHITE
    
    for c_idx, (col_name, width, default_align) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row_idx, column=c_idx)
        val = tc_data.get(col_name, "")
        cell.value = val
        cell.border = CELL_BORDER
        cell.fill = base_fill
        cell.alignment = default_align
        cell.font = FONT_DATA
        
        # Specific formatting per column
        if col_name == "TC ID":
            cell.font = FONT_DATA_BOLD
            cell.alignment = ALIGN_CENTER
        elif col_name == "Execution Status":
            if val == "PASSED":
                cell.fill = FILL_PASSED
                cell.font = FONT_PASSED
            elif "OBS" in val or "OBSERVATION" in val:
                cell.fill = FILL_OBS
                cell.font = FONT_OBS
            elif "DEFERRED" in val:
                cell.fill = FILL_DEFERRED
                cell.font = FONT_DEFERRED
        elif col_name == "Severity / Priority":
            if "P0" in val or "Critical" in val:
                cell.font = FONT_P0
            elif "P1" in val or "High" in val:
                cell.font = FONT_P1
            elif "P2" in val or "Medium" in val:
                cell.font = FONT_P2
            else:
                cell.font = FONT_P3
        elif col_name == "Defect Ref" and val != "-" and val != "None":
            cell.font = FONT_P1

def create_kpi_card(ws, start_row, start_col, title, value, subtext=None, val_color="1E3A8A"):
    # Merges 2 rows x 2 cols
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row+1, end_column=start_col+1)
    cell = ws.cell(row=start_row, column=start_col)
    cell.value = f"{value}\n{title}"
    cell.font = Font(name="Segoe UI", size=13, bold=True, color=val_color)
    cell.fill = FILL_KPI_CARD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for r in range(start_row, start_row+2):
        for c in range(start_col, start_col+2):
            c_cell = ws.cell(row=r, column=c)
            c_cell.border = KPI_BORDER
            c_cell.fill = FILL_KPI_CARD

print("Loaded base openpyxl script modules.")
