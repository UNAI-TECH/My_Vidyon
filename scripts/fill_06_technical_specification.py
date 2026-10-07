import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import shutil

def fill_cell(cell, text, bold=False, italic=False, size_pt=9.5, color_rgb=(30, 41, 59), bg_hex=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(*color_rgb)
    if bg_hex:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

def add_paragraph_to_cell(cell, text, bold=False, italic=False, size_pt=9.5, color_rgb=(30, 41, 59), space_after=3):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(*color_rgb)
    return p

def process_tech_spec(src_path, dest_paths):
    doc = docx.Document(src_path)
    
    # -------------------------------------------------------------
    # Table 0: Document Control (4 rows, 6 cols)
    # -------------------------------------------------------------
    t0 = doc.tables[0]
    t0_headers = ['Version', 'Date', 'Prepared By', 'Reviewed By', 'Approved By', 'Description of Changes']
    for c_idx, h in enumerate(t0_headers):
        fill_cell(t0.rows[0].cells[c_idx], h, bold=True, size_pt=9.5, color_rgb=(255, 255, 255), bg_hex="1E293B", align=WD_ALIGN_PARAGRAPH.CENTER)
    
    t0_data = [
        ['1.0', '2026-03-05', 'Kamalesh S', 'Kamalesh S', 'Kamalesh S', 'Initial Technical Specification baseline for My Vidyon ERP'],
        ['1.1', '2026-03-20', 'Kamalesh S', 'Kamalesh S', 'Kamalesh S', 'Integration of Supabase Edge Functions, LargeSecureStore, and FCM push services'],
        ['1.2', '2026-10-07', 'Kamalesh S', 'Kamalesh S', 'Kamalesh S', 'Full-Stack Rebuild: Expo SDK 57, React Native 0.86, Electron Windows build, 33-module RBAC']
    ]
    for r_idx, row_vals in enumerate(t0_data, start=1):
        for c_idx, val in enumerate(row_vals):
            bg = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 2, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(t0.rows[r_idx].cells[c_idx], val, size_pt=9, bg_hex=bg, align=align)

    # -------------------------------------------------------------
    # Table 1: Project Metadata (2 rows, 2 cols)
    # -------------------------------------------------------------
    t1 = doc.tables[1]
    fill_cell(t1.rows[0].cells[0], "Project Name: My Vidyon ERP", bold=True, size_pt=9.5, bg_hex="F1F5F9")
    fill_cell(t1.rows[0].cells[1], "Module / Component: Full-Stack Multi-Tenant System", bold=True, size_pt=9.5, bg_hex="F1F5F9")
    fill_cell(t1.rows[1].cells[0], "Prepared By: Kamalesh S (Lead Architect & CTO)", size_pt=9.5, bg_hex="FFFFFF")
    fill_cell(t1.rows[1].cells[1], "Date: October 7, 2026", size_pt=9.5, bg_hex="FFFFFF")

    # -------------------------------------------------------------
    # Table 2: 1. OVERVIEW (1 row, 1 col)
    # -------------------------------------------------------------
    t2 = doc.tables[2]
    c2 = t2.rows[0].cells[0]
    overview_text_1 = (
        "My Vidyon ERP is a unified, cloud-native K-12 Educational Resource Planning platform engineered "
        "to deliver end-to-end digitization for educational institutions. The platform is designed with a "
        "modern cross-platform client architecture using React Native / Expo (SDK 57) that targets three "
        "distinct presentation layers: Android Native (via EAS APK/AAB), Web (via React Native Web static export), "
        "and Windows Desktop (via an Electron Win32 x64 wrapper running a local HTTP distribution server)."
    )
    overview_text_2 = (
        "The technical architecture couples client-side reactive state management via TanStack React Query v5 "
        "with Supabase Cloud Backend-as-a-Service (PostgreSQL 15+, GoTrue Auth, Realtime WebSockets, and Deno Edge Functions). "
        "The system enforces strict multi-tenant Row-Level Security (RLS), hardware-backed JWT keychain encryption "
        "via LargeSecureStore, a centralized 33-module client/server RBAC permission matrix (can() evaluator), and 10 "
        "asynchronous Supabase Realtime channels providing sub-second cache invalidation across all connected role dashboards."
    )
    fill_cell(c2, overview_text_1, size_pt=9.5, color_rgb=(15, 23, 42), bg_hex="FFFFFF")
    add_paragraph_to_cell(c2, overview_text_2, size_pt=9.5, color_rgb=(15, 23, 42))

    # -------------------------------------------------------------
    # Table 3: 2. TECHNOLOGY STACK (6 rows, 4 cols)
    # -------------------------------------------------------------
    t3 = doc.tables[3]
    t3_headers = ['Layer', 'Technology', 'Version', 'Notes']
    for c_idx, h in enumerate(t3_headers):
        fill_cell(t3.rows[0].cells[c_idx], h, bold=True, size_pt=9.5, color_rgb=(255, 255, 255), bg_hex="1E293B", align=WD_ALIGN_PARAGRAPH.CENTER)

    t3_data = [
        ['Frontend Core', 'React Native / Expo', 'SDK 57 (~57.0.25, RN 0.86.3)', 'Expo Router v57 file-based typed routing, Lucide icons, Reanimated 4.5.1'],
        ['Desktop Runtime', 'Electron (Win32 x64)', 'Standalone Packager', 'scripts/build-exe.mjs local HTTP runner serving static web export (dist/)'],
        ['Backend & Database', 'Supabase (PostgreSQL)', 'PostgreSQL 15.x', '52+ relational tables, Row-Level Security (RLS), GoTrue Auth, 80+ migrations'],
        ['State & Caching', 'TanStack React Query', 'v5.90.21', 'Optimistic mutations, persistent server-state cache, invalidation via Realtime'],
        ['Serverless & Realtime', 'Supabase Edge / Realtime', 'Deno Runtime / WebSockets', 'Edge Functions (create-user, delete-user, push dispatch), 10 Realtime channels']
    ]
    for r_idx, row_vals in enumerate(t3_data, start=1):
        for c_idx, val in enumerate(row_vals):
            bg = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(t3.rows[r_idx].cells[c_idx], val, size_pt=9, bg_hex=bg, align=align)

    # -------------------------------------------------------------
    # Table 4: 3. SYSTEM COMPONENTS (6 rows, 3 cols)
    # -------------------------------------------------------------
    t4 = doc.tables[4]
    t4_headers = ['Component', 'Responsibility', 'Dependencies']
    for c_idx, h in enumerate(t4_headers):
        fill_cell(t4.rows[0].cells[c_idx], h, bold=True, size_pt=9.5, color_rgb=(255, 255, 255), bg_hex="1E293B", align=WD_ALIGN_PARAGRAPH.CENTER)

    t4_data = [
        ['Auth & Session Manager (useAuth.tsx)', 'Session lifecycle, token chunking via LargeSecureStore, role stability locks, blocked account booting', 'Supabase GoTrue, expo-secure-store, AsyncStorage'],
        ['Central RBAC Evaluator (src/lib/rbac.ts)', 'Evaluates 8 actions across 33 modules via can(), enforces BASELINE_ROLE_PERMISSIONS matrix and stakeholder lock', 'Role enum, PermissionCheckContext, RBACContext'],
        ['Real-Time Sync Bridge (useERPRealtime.ts)', 'Subscribes to 10 Supabase channels, auto-invalidates React Query caches on Postgres INSERT/UPDATE/DELETE', 'Supabase Realtime, QueryClient, Expo Notifications'],
        ['Academic & Timetable Engine', 'Manages class standards, subjects, faculty schedules, and collision-free timetable allocation', 'PostgreSQL timetable_slots, classes, useInstitutionTimetable'],
        ['Financial & Fee Engine (useFeeWorkflow.ts)', 'Configures multi-component fee structures, records cash bills, generates PDF receipts, handles online payments', 'student_fees, fee_payments, invoiceGenerator.ts, expo-print']
    ]
    for r_idx, row_vals in enumerate(t4_data, start=1):
        for c_idx, val in enumerate(row_vals):
            bg = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
            fill_cell(t4.rows[r_idx].cells[c_idx], val, size_pt=9, bg_hex=bg)

    # -------------------------------------------------------------
    # Table 5: 4. DATA MODEL (1 row, 1 col)
    # -------------------------------------------------------------
    t5 = doc.tables[5]
    c5 = t5.rows[0].cells[0]
    dm_p1 = (
        "The My Vidyon data architecture is organized into 52+ normalized PostgreSQL tables partitioned "
        "across 7 functional domains with relational integrity and Row-Level Security:\n"
        "• Core & Auth: 'profiles' (user roles, institution mapping, status), 'user_sessions' (device and IP tracking), "
        "'user_push_tokens' (FCM mobile push registration), 'platform_activities' (system audit trail).\n"
        "• Institutional Hierarchy: 'institutions' (tenants, domain settings, branding), 'classes' (grade standards and sections), "
        "'subjects' (master curriculum), 'faculty_subjects' (teacher subject allocations).\n"
        "• Staff & Academics: 'staff_details' (payroll metadata, departments), 'staff_attendance' (daily biometric/manual logs), "
        "'leave_requests' / 'staff_leaves' (faculty leave applications with approval workflows).\n"
        "• Students & Guardians: 'students' (roll number, personal details, photos), 'parents' (guardian profiles), "
        "'parent_student_relations' (multi-child linkage), 'student_attendance' (period/daily attendance records).\n"
        "• Examinations & Grades: 'timetable' / 'timetable_slots' (weekly schedules), 'exams' / 'exam_schedules' (test dates), "
        "'exam_results' / 'grades' (draft/published marks), 'assignments' / 'submissions' (homework workflows).\n"
        "• Finance & Commercial: 'fee_structures' (tier-wise breakdown), 'student_fees' (ledger balances), "
        "'fee_payments' (transaction logs and receipts), 'ad_leads' / 'subscriptions' (monetization).\n"
        "• Canteen & System: 'canteen_sessions' (meal schedules), 'canteen_attendance' (student dining entries), "
        "'announcements' (broadcasts), 'notifications' (in-app alerts), plus 10 dedicated Supabase Storage Buckets."
    )
    fill_cell(c5, dm_p1, size_pt=9.5, color_rgb=(15, 23, 42), bg_hex="FFFFFF")

    # -------------------------------------------------------------
    # Table 6: 5. API SPECIFICATIONS (6 rows, 4 cols)
    # -------------------------------------------------------------
    t6 = doc.tables[6]
    t6_headers = ['Endpoint / Service', 'Method', 'Description', 'Request / Response Contract']
    for c_idx, h in enumerate(t6_headers):
        fill_cell(t6.rows[0].cells[c_idx], h, bold=True, size_pt=9.5, color_rgb=(255, 255, 255), bg_hex="1E293B", align=WD_ALIGN_PARAGRAPH.CENTER)

    t6_data = [
        ['/auth/v1/token?grant_type=password', 'POST', 'Supabase GoTrue User Authentication', 'Req: {email, password}\nRes: {access_token, refresh_token, user: {id, email}}'],
        ['/rest/v1/profiles?id=eq.{userId}', 'GET', 'Fetch authenticated profile and effective role', 'Req: Bearer JWT\nRes: [{role, institution_id, full_name, is_active, ...}]'],
        ['/functions/v1/create-user', 'POST', 'Edge Function admin user provisioning', 'Req: {email, password, role, institution_id, full_name}\nRes: {user_id, success: true}'],
        ['/functions/v1/send-push-notification', 'POST', 'FCM Push notification dispatcher', 'Req: {user_id, title, body, data: {action_url}}\nRes: {status: "sent", ticket_id}'],
        ['/rest/v1/student_attendance', 'POST', 'Batch record daily student attendance', 'Req: [{student_id, class_id, date, status: "present"|"absent"}]\nRes: 201 Created HTTP status']
    ]
    for r_idx, row_vals in enumerate(t6_data, start=1):
        for c_idx, val in enumerate(row_vals):
            bg = "F8F9FA" if r_idx % 2 == 1 else "FFFFFF"
            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(t6.rows[r_idx].cells[c_idx], val, size_pt=8.5, bg_hex=bg, align=align)

    # -------------------------------------------------------------
    # Section 6: Security Considerations (Paragraphs after Table 6)
    # -------------------------------------------------------------
    sec_p = doc.paragraphs[15] if len(doc.paragraphs) > 15 else doc.add_paragraph()
    sec_p.text = (
        "6.1 Multi-Tenant Isolation: Enforced at the database layer via PostgreSQL Row-Level Security (RLS). "
        "All tables scope row access to the authenticated user's institution_id.\n"
        "6.2 Token Protection: On Android/iOS, access/refresh tokens are stored in hardware-backed encrypted "
        "KeyStore partitions via LargeSecureStore with chunking to prevent 2048-byte overflow. On web/desktop, "
        "tokens reside in local isolated application storage.\n"
        "6.3 Input Sanitization: PostgREST parameterized queries eliminate SQL injection vectors. React Native "
        "natively escapes JSX children to prevent XSS script execution.\n"
        "6.4 Real-Time Access Eviction: Profiles with is_active = false or suspended institutions are immediately "
        "detected by useAuth and forcibly logged out.\n"
        "6.5 View-Only Enforcement: The institution_stakeholder role is governed by database write-deny triggers and "
        "client-side PermissionGate wrappers prohibiting mutations across all 33 ERP modules."
    )
    for run in sec_p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(15, 23, 42)

    # -------------------------------------------------------------
    # Table 7: APPROVAL (1 row, 2 cols)
    # -------------------------------------------------------------
    t7 = doc.tables[7]
    fill_cell(t7.rows[0].cells[0], "Prepared By Name: Kamalesh S (Lead Architect & CTO)\nDate: October 7, 2026\nSignature: [Signed Electronically]", size_pt=9.5, bg_hex="F8F9FA")
    fill_cell(t7.rows[0].cells[1], "Approved By Name: Kamalesh S (Technical Lead)\nDate: October 7, 2026\nSignature: [Approved for Production Baseline]", size_pt=9.5, bg_hex="F8F9FA")

    # Save to all destination locations
    for d in dest_paths:
        os.makedirs(os.path.dirname(os.path.abspath(d)), exist_ok=True)
        doc.save(d)
        print(f"Saved Technical Specification to: {d}")

if __name__ == '__main__':
    src = os.path.join('docs', 'ai', 'documentations', 'Copy of 06_Technical_Specification.docx')
    destinations = [
        os.path.join('docs', 'ai', 'documentations', 'Copy of 06_Technical_Specification.docx'),
        os.path.join('docs', 'ai', 'documentations', 'docs_filled', 'Copy of 06_Technical_Specification.docx'),
        os.path.join('docs_filled', 'Copy of 06_Technical_Specification.docx')
    ]
    process_tech_spec(src, destinations)
