import docx
import os

docs_to_inspect = [
    'Copy of 06_Technical_Specification.docx',
    'Copy of 07_Architecture_Document.docx',
    'Copy of 09_Sprint_Plan.docx',
    'Copy of 10_Test_Plan.docx'
]

for doc_name in docs_to_inspect:
    path = os.path.join('docs', 'ai', 'documentations', doc_name)
    doc = docx.Document(path)
    print(f'==============================')
    print(f'FILE: {doc_name}')
    print(f'==============================')
    print('PARAGRAPHS:')
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f'  P{i}: "{p.text}"')
    print('TABLES:')
    for t_idx, table in enumerate(doc.tables):
        print(f'  Table {t_idx} (rows: {len(table.rows)}, cols: {len(table.columns)}):')
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            print(f'    R{r_idx}: {cells}')
